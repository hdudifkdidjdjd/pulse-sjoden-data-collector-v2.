
"""Pulse Sjöden V85 Excel Import v0.2.

Converts reviewed Excel exports to V85 CSV files.
Column names must be mapped explicitly.
Does not modify the existing database or SAT engine.
"""

import argparse
import csv
import json
from datetime import date as Date
from pathlib import Path
from urllib.parse import urlsplit


START_FIELDS = (
    "date", "track", "division", "start_number",
    "horse", "driver", "scratched", "source_url"
)

RESULT_FIELDS = (
    "date", "track", "division",
    "winning_number", "source_url"
)

ALLOWED_HOSTS = {
    "travsport.se",
    "www.travsport.se",
    "atg.se",
    "www.atg.se",
}


def safe_source(url):
    try:
        p = urlsplit(url)
        return (
            p.scheme == "https"
            and p.hostname in ALLOWED_HOSTS
            and not p.username
            and not p.password
            and p.port in (None, 443)
        )
    except ValueError:
        return False


def prepare_rows(
    headers,
    records,
    mapping,
    kind,
    date,
    track,
    division,
    source_url,
):
    if kind == "starts":
        required = (
            "start_number",
            "horse",
            "driver",
        )
    else:
        required = (
            "start_number",
            "finish_position",
        )

    if not all(
        mapping.get(key) in headers
        for key in required
    ):
        raise ValueError(
            "Missing mapped source columns: "
            + ", ".join(required)
        )

    if not 1 <= division <= 8:
        raise ValueError(
            "V85 division must be 1-8"
        )

    if not safe_source(source_url):
        raise ValueError(
            "Unapproved source URL"
        )

    if not date or not track:
        raise ValueError(
            "Date and track are required"
        )

    output = []
    seen = set()

    for line_number, row in enumerate(records, 2):
        if all(
            value is None
            or str(value).strip() == ""
            for value in row.values()
        ):
            continue

        try:
            number = int(
                row[mapping["start_number"]]
            )

            if number < 1 or number in seen:
                raise ValueError(
                    "Invalid or duplicate starting number"
                )

            seen.add(number)

            if kind == "starts":
                horse = str(
                    row[mapping["horse"]] or ""
                ).strip()

                driver = str(
                    row[mapping["driver"]] or ""
                ).strip()

                if not horse or not driver:
                    raise ValueError(
                        "Missing horse or driver"
                    )

                raw = row.get(
                    mapping.get("scratched", ""),
                    0,
                )

                if (
                    raw is None
                    or str(raw).strip() == ""
                ):
                    raw = "0"
                else:
                    raw = str(raw).strip()

                if raw not in ("0", "1"):
                    raise ValueError(
                        "Scratched must be 0 or 1"
                    )

                output.append((
                    date,
                    track,
                    division,
                    number,
                    horse,
                    driver,
                    raw,
                    source_url,
                ))

            else:
                position = str(
                    row[mapping["finish_position"]]
                    or ""
                ).strip()

                if position in ("1", "1."):
                    output.append((
                        date,
                        track,
                        division,
                        number,
                        source_url,
                    ))

        except (
            TypeError,
            ValueError,
            KeyError,
        ) as error:
            raise ValueError(
                f"Spreadsheet row {line_number}: {error}"
            ) from error

    if kind == "results" and len(output) != 1:
        raise ValueError(
            "Expected exactly one winner, "
            f"found {len(output)}"
        )

    if kind == "starts" and not output:
        raise ValueError(
            "No starting horses found"
        )

    return output


def convert(
    path,
    mapping_path,
    output,
    kind,
    date,
    track,
    division,
    source_url,
    sheet=None,
):
    from openpyxl import load_workbook

    if kind not in ("starts", "results"):
        raise ValueError(
            "Kind must be starts or results"
        )

    Date.fromisoformat(date)

    mapping = json.loads(
        Path(mapping_path).read_text(
            encoding="utf-8"
        )
    )

    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    try:
        ws = (
            workbook[sheet]
            if sheet
            else workbook.active
        )

        iterator = ws.iter_rows(
            values_only=True
        )

        try:
            headers = [
                str(value or "").strip()
                for value in next(iterator)
            ]
        except StopIteration as error:
            raise ValueError(
                "Empty spreadsheet"
            ) from error

        if (
            len(headers) != len(set(headers))
            or not all(headers)
        ):
            raise ValueError(
                "Empty or duplicate Excel headers"
            )

        records = (
            dict(zip(headers, values))
            for values in iterator
        )

        rows = prepare_rows(
            headers,
            records,
            mapping,
            kind,
            date,
            track,
            division,
            source_url,
        )

    finally:
        workbook.close()

    output = Path(output)
    output.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with output.open(
        "w",
        newline="",
        encoding="utf-8",
    ) as stream:
        writer = csv.writer(stream)

        writer.writerow(
            START_FIELDS
            if kind == "starts"
            else RESULT_FIELDS
        )

        writer.writerows(rows)

    return len(rows)


def main():
    parser = argparse.ArgumentParser(
        description=__doc__
    )

    parser.add_argument(
        "--xlsx",
        required=True,
    )

    parser.add_argument(
        "--mapping",
        required=True,
    )

    parser.add_argument(
        "--output",
        required=True,
    )

    parser.add_argument(
        "--kind",
        required=True,
        choices=("starts", "results"),
    )

    parser.add_argument(
        "--date",
        required=True,
    )

    parser.add_argument(
        "--track",
        required=True,
    )

    parser.add_argument(
        "--division",
        required=True,
        type=int,
    )

    parser.add_argument(
        "--source-url",
        required=True,
    )

    parser.add_argument("--sheet")

    args = parser.parse_args()

    count = convert(
        args.xlsx,
        args.mapping,
        args.output,
        args.kind,
        args.date,
        args.track,
        args.division,
        args.source_url,
        args.sheet,
    )

    print(json.dumps({
        "kind": args.kind,
        "rows": count,
        "output": args.output,
    }))


if __name__ == "__main__":
    main()
